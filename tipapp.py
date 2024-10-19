from flask import Flask, render_template, request, redirect, url_for, send_file
from docx import Document
from datetime import datetime
import os
import sys
import sqlite3
import pandas as pd
from io import BytesIO

app = Flask(__name__)

# Database configuration
DATABASE = 'tippani.db'

# Initialize the database
def init_db():
    conn = sqlite3.connect(DATABASE)
    cursor = conn.cursor()
    cursor.execute('''
        CREATE TABLE IF NOT EXISTS submissions (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            subject TEXT,
            input1 TEXT NOT NULL,
            input2 TEXT,
            input3 TEXT,
            input4 TEXT,
            input5 TEXT,
            input6 TEXT,
            input7 TEXT,
            input8 TEXT,
            input9 TEXT,
            input10 TEXT,
            input11 TEXT,
            input12 TEXT,
            input13 TEXT,
            input14 TEXT,
            input15 TEXT
        )
    ''')
    conn.commit()
    conn.close()

# Call this to initialize the DB when the app starts
init_db()

# Template dictionary
templates = {
    "template1": "विषय १",
    "template2": "विषय २",
    "template3": "विषय ३",
    "template4": "विषय ४",
    "template5": "विषय ५",
    "template6": "विषय ६",
    "template7": "विषय ७",
    "template8": "विषय ८",
    "template9": "विषय ९",
    "template10": "विषय १०",
}

# Utility function to handle paths for both PyInstaller and normal environments
def resource_path(relative_path):
    """ Get the absolute path to the resource, works for both development and PyInstaller """
    try:
        base_path = sys._MEIPASS  # PyInstaller temporary folder
    except AttributeError:
        base_path = os.path.abspath(".")
    return os.path.join(base_path, relative_path)

@app.route('/')
def index():
    return render_template('index.html', templates=templates)

@app.route('/select_template', methods=['POST'])
def select_template():
    selected_template = request.form.get('subject')
    return redirect(url_for('show_template', template=selected_template))

@app.route('/template/<template>')
def show_template(template):
    if template in templates:
        # Fetch the most recent submission for this template
        conn = sqlite3.connect(DATABASE)
        cursor = conn.cursor()
        cursor.execute('SELECT * FROM submissions WHERE subject = ? ORDER BY id DESC LIMIT 1', (template,))
        previous_data = cursor.fetchone()
        conn.close()

        # If there's previous data, fill the form with it, otherwise leave it blank
        if previous_data:
            default_values = {f'input{i}': previous_data[i+2] if len(previous_data) > i+2 else '' for i in range(1, 16)}
        else:
            default_values = {f'input{i}': '' for i in range(1, 16)}

        return render_template(f'{template}.html', default_values=default_values)
    
    return "Template not found", 404


@app.route('/generate', methods=['POST'])
def generate_tippani():
    selected_subject = request.form.get('subject')
    template_file = f"{selected_subject}.docx"

    # Collect form inputs, allowing null values
    inputs = {f'input{i}': request.form.get(f'input{i}') or '' for i in range(1, 16)}

    if template_file and inputs:
        # Load the selected DOCX template
        doc = Document(resource_path(f'templates/{template_file}'))

        # Replace placeholders in the document with user inputs
        for i in range(1, 16):
            placeholder = f'<{i}>'
            value = inputs[f'input{i}']
            for paragraph in doc.paragraphs:
                if placeholder in paragraph.text:
                    paragraph.text = paragraph.text.replace(placeholder, value)

        # Create a unique filename with a timestamp
        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')  # Format: YYYYMMDD_HHMMSS
        output_filename = f'tippani_{selected_subject}_{timestamp}.docx'
        output_path = os.path.join('output', output_filename)

        # Save the document
        if not os.path.exists('output'):
            os.makedirs('output')
        doc.save(output_path)

        # Insert the submission into the database
        conn = sqlite3.connect(DATABASE)
        cursor = conn.cursor()
        cursor.execute('''
            INSERT INTO submissions (subject, input1, input2, input3, input4, input5, input6, input7, input8, input9, input10, input11, input12, input13, input14, input15)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        ''', (selected_subject, *inputs.values()))
        conn.commit()
        conn.close()

        return f"Tippani generated successfully! Download here: {output_path}"

    return "Error in generating Tippani."

@app.route('/download_data')
def download_data():
    # Fetch all data from the database
    conn = sqlite3.connect(DATABASE)
    query = "SELECT * FROM submissions"
    df = pd.read_sql_query(query, conn)
    conn.close()

    # Convert the data to an Excel file and send it as a downloadable file
    output = BytesIO()
    writer = pd.ExcelWriter(output, engine='xlsxwriter')
    df.to_excel(writer, index=False, sheet_name='Submissions')
    writer.close()
    output.seek(0)

    return send_file(output, download_name='submissions_data.xlsx', as_attachment=True)

if __name__ == '__main__':
    app.run(debug=True)
